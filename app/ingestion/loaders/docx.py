from pathlib import Path

from docx import Document as DocxDocument

from app.ingestion.loaders.base import DocumentLoader, LoadedPage
from app.ingestion.structure import TableData
from app.ingestion.table_serializer import table_to_markdown


class DocxLoader(DocumentLoader):
    """Load Microsoft Word (.docx) documents as a single logical page.

    Extracts body paragraph text in document order, skipping empty
    paragraphs, with paragraph boundaries preserved as newlines. Word
    documents don't carry reliable page-break metadata without rendering,
    so (like MarkdownLoader) no attempt is made to infer physical page
    numbers -- the whole document is returned as one LoadedPage.

    If the document contains tables, each one's genuinely structured
    `document.tables` data (python-docx already exposes real rows/cells,
    unlike PDF text extraction) is rendered as a standard Markdown pipe
    table -- via the same `table_serializer.table_to_markdown` used
    elsewhere in ingestion -- and appended after all paragraph text, rather
    than interleaved at each table's true position. python-docx exposes
    `document.paragraphs` and `document.tables` as two separate flat
    sequences; reconstructing their true interleaved reading order would
    require walking the raw document XML, which is unnecessary complexity
    for this loader. Rendering as real pipe-table syntax (header row +
    `|---|---|` separator), rather than the previous flat
    `"cell | cell"`-per-row text, is what lets `StructureExtractor` detect
    these as genuine TABLE structural units downstream -- the first row of
    each table is treated as its header row.
    """

    def load(self, path: Path) -> list[LoadedPage]:
        document = DocxDocument(path)

        paragraphs: list[str] = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue

            style_name = paragraph.style.name if paragraph.style is not None else ""

            if style_name.startswith("Heading "):
                try:
                    level = int(style_name.removeprefix("Heading ").strip())
                except ValueError:
                    level = 0

                if 1 <= level <= 6:
                    text = f"{'#' * level} {text}"

            paragraphs.append(text)

        table_sections: list[str] = []

        for table_index, table in enumerate(document.tables, start=1):
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            rows = [row for row in rows if any(cell for cell in row)]

            if not rows:
                continue

            headers, *data_rows = rows

            table_data = TableData(
                table_id=f"docx-table-{table_index}",
                # .docx has no reliable "this is the title" signal separate
                # from a preceding paragraph -- never fabricate one.
                title=None,
                headers=headers,
                rows=data_rows,
                page_numbers=[1],
            )

            table_sections.append(table_to_markdown(table_data))

        sections = ["\n".join(paragraphs), *table_sections]

        # Blank-line-separated so StructureExtractor always sees a clean
        # paragraph/table boundary, regardless of what the last paragraph
        # line looked like.
        content = "\n\n".join(section for section in sections if section)

        return [LoadedPage(page_number=1, content=content)]
