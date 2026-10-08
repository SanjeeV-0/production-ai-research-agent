from pathlib import Path

from docx import Document

OUTPUT_DIR = Path(__file__).parent / "benchmark"


def add_paragraphs(document: Document, paragraphs: list[str]) -> None:
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)


def add_section(
    document: Document,
    heading: str,
    paragraphs: list[str],
    level: int = 1,
) -> None:
    document.add_heading(heading, level=level)
    add_paragraphs(document, paragraphs)


def add_table(
    document: Document,
    headers: list[str],
    rows: list[list[str]],
) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"

    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header

    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value


def repeated_paragraphs(
    topic: str,
    concepts: list[str],
    count: int,
) -> list[str]:
    paragraphs: list[str] = []

    templates = [
        (
            "{topic} depends on {concept}. In a production retrieval system, "
            "the behavior of {concept} should be evaluated together with "
            "candidate generation, ranking, and evidence quality. "
            "The implementation should preserve enough context to distinguish "
            "a genuinely useful match from text that is merely similar in wording."
        ),
        (
            "A practical implementation of {concept} should preserve enough "
            "context for downstream answer generation. The same mechanism can "
            "behave differently when documents contain overlapping terminology "
            "or conflicting measurements. Evaluation should therefore examine "
            "both retrieval quality and the evidence actually returned."
        ),
        (
            "Evaluation of {concept} should distinguish retrieval failure from "
            "ranking failure. If relevant evidence is absent from the candidate "
            "pool, a later reranking stage cannot recover it. Candidate coverage "
            "and final ordering should therefore be measured as separate stages."
        ),
        (
            "Operational measurements for {concept} should include both quality "
            "and efficiency. Recall, precision, latency, and the number of "
            "candidates processed are useful complementary signals. A change "
            "that improves one metric can still introduce an undesirable "
            "trade-off elsewhere in the retrieval pipeline."
        ),
        (
            "When investigating {concept}, engineers should inspect the boundary "
            "between indexing, candidate generation, and ranking. Similar "
            "terminology may occur in several sections while the underlying "
            "claims differ. Retrieval evaluation should preserve enough source "
            "context to make those distinctions visible."
        ),
        (
            "A retrieval system using {concept} should be evaluated under "
            "different query formulations rather than a single wording. "
            "Paraphrases, related terminology, and multi-part questions can "
            "exercise different retrieval behaviors. Stable evaluation cases "
            "make those differences easier to diagnose."
        ),
    ]

    for index in range(count):
        paragraphs.append(
            templates[index % len(templates)].format(
                topic=topic,
                concept=concepts[index % len(concepts)],
            )
        )

    return paragraphs


def create_vector_search_fundamentals() -> None:
    document = Document()

    add_section(
        document,
        "Vector Search Fundamentals",
        [
            "Vector search retrieves information by comparing numerical "
            "representations of queries and documents. Semantic embeddings "
            "allow a retrieval system to match concepts even when exact "
            "keywords differ.",
            "The benchmark deliberately uses related terminology such as "
            "semantic retrieval, dense retrieval, embedding search, and "
            "vector similarity. These terms are related but should not be "
            "treated as interchangeable evidence in every query.",
        ],
    )

    add_section(
        document,
        "Embeddings",
        [
            "An embedding model maps text into a fixed-dimensional vector "
            "space. The dimensionality is determined by the selected model "
            "and must remain compatible between indexing and querying.",
            "Raw model outputs may be stored without normalization. Cosine "
            "distance can then be used directly by the database to compare "
            "query and document vectors.",
        ],
    )

    add_section(
        document,
        "Similarity Search",
        [
            "Cosine similarity measures the angular relationship between "
            "vectors. When cosine distance is used, lower values represent "
            "closer vectors.",
            "A candidate limit controls how many results are initially "
            "returned. This limit affects recall, processing cost, and the "
            "amount of material presented to later ranking stages.",
        ],
    )

    add_section(
        document,
        "Chunking",
        repeated_paragraphs(
            "Chunking",
            [
                "semantic boundaries",
                "paragraph boundaries",
                "section metadata",
                "retrieval context",
            ],
            220,
        ),
    )

    add_section(
        document,
        "Recall and Precision",
        repeated_paragraphs(
            "Retrieval evaluation",
            [
                "recall",
                "precision",
                "candidate coverage",
                "top-k ranking",
            ],
            180,
        ),
    )

    add_section(
        document,
        "Controlled Benchmark Observation",
        [
            "This document uses a candidate pool of 50 as an example "
            "configuration that is sufficient for many ordinary workloads.",
            "That recommendation is intentionally not universal. Other "
            "documents in the benchmark recommend larger candidate pools "
            "for high-recall workloads.",
        ],
    )

    document.save(OUTPUT_DIR / "vector_search_fundamentals.docx")


def create_hnsw_document() -> None:
    document = Document()

    add_section(
        document,
        "HNSW Indexing Deep Dive",
        [
            "Hierarchical Navigable Small World indexing provides an "
            "approximate nearest-neighbor search structure. It organizes "
            "vectors into graph layers so that search can navigate from "
            "coarse regions toward increasingly local neighborhoods.",
            "HNSW parameters influence the trade-off between index size, "
            "construction cost, search effort, latency, and recall.",
        ],
    )

    add_section(
        document,
        "Graph Construction",
        repeated_paragraphs(
            "HNSW graph construction",
            [
                "graph connectivity",
                "neighbor selection",
                "construction effort",
                "layer assignment",
            ],
            210,
        ),
    )

    add_section(
        document,
        "Search Parameters",
        repeated_paragraphs(
            "HNSW search",
            [
                "ef search",
                "candidate exploration",
                "nearest-neighbor traversal",
                "recall under approximate search",
            ],
            190,
        ),
    )

    add_section(
        document,
        "Candidate Pool Recommendations",
        [
            "For high-recall retrieval workloads, this document recommends "
            "testing a candidate pool of 100 rather than assuming that 50 "
            "candidates are sufficient.",
            "This is intentionally a controlled disagreement with the "
            "Vector Search Fundamentals document. The benchmark should "
            "reward retrieval that identifies the relevant evidence rather "
            "than assuming that one configuration is universally correct.",
        ],
    )

    add_section(
        document,
        "Operational Trade-offs",
        repeated_paragraphs(
            "HNSW operations",
            [
                "index construction",
                "memory consumption",
                "query latency",
                "recall measurement",
            ],
            180,
        ),
    )

    rows = [
        [
            f"configuration-{i:03d}",
            str(384 + (i % 4) * 128),
            str(20 + (i % 6) * 10),
            str(50 + (i % 8) * 10),
            f"{0.78 + (i % 20) / 100:.2f}",
        ]
        for i in range(1, 181)
    ]

    add_section(
        document,
        "HNSW Configuration Benchmark Table",
        [
            "The following table records controlled benchmark configurations. "
            "The table is intentionally large so that ingestion must fragment "
            "it into multiple retrieval chunks while preserving its header.",
        ],
    )

    add_table(
        document,
        ["Configuration", "Dimension", "M", "Candidate Pool", "Recall"],
        rows,
    )

    document.save(OUTPUT_DIR / "hnsw_indexing_deep_dive.docx")


def create_reranking_document() -> None:
    document = Document()

    add_section(
        document,
        "Reranking with Cross-Encoders",
        [
            "A reranker receives an initial candidate set and scores each "
            "candidate against the original query. Cross-encoders can model "
            "query-document interactions more directly than independent "
            "embedding comparisons.",
            "Reranking cannot recover evidence that was excluded during "
            "candidate generation. Candidate recall therefore remains an "
            "important diagnostic.",
        ],
    )

    add_section(
        document,
        "Candidate Generation",
        repeated_paragraphs(
            "Candidate generation",
            [
                "vector similarity",
                "candidate recall",
                "query representation",
                "candidate pool size",
            ],
            210,
        ),
    )

    add_section(
        document,
        "Cross-Encoder Scoring",
        repeated_paragraphs(
            "Cross-encoder reranking",
            [
                "query-document interaction",
                "relevance scoring",
                "ranking precision",
                "reranking latency",
            ],
            190,
        ),
    )

    add_section(
        document,
        "Reranking Evaluation",
        repeated_paragraphs(
            "Reranking evaluation",
            [
                "candidate recall",
                "final recall",
                "precision at three",
                "precision at five",
            ],
            65,
        ),
    )

    rows = [
        [
            f"model-{i:03d}",
            f"query-{(i % 20) + 1:02d}",
            str(20 + (i % 5) * 20),
            f"{0.50 + (i % 45) / 100:.2f}",
            f"{0.45 + (i % 50) / 100:.2f}",
            f"{12 + (i % 30)}",
        ]
        for i in range(1, 161)
    ]

    add_section(
        document,
        "Reranking Results Table",
        [
            "This table intentionally contains several related ranking "
            "configurations. Some configurations have similar scores but "
            "different latency characteristics.",
        ],
    )

    add_table(
        document,
        [
            "Model",
            "Query",
            "Candidates",
            "Candidate Recall",
            "Final Precision",
            "Latency ms",
        ],
        rows,
    )

    add_section(
        document,
        "Known Limitation",
        [
            "A reranker can improve ordering among retrieved candidates, "
            "but it cannot compensate for an insufficient candidate pool.",
            "This distinction is important when interpreting evaluation "
            "results because a strong final ranking score can conceal a "
            "candidate-generation problem.",
        ],
    )

    document.save(OUTPUT_DIR / "reranking_with_cross_encoders.docx")


def create_query_decomposition_document() -> None:
    document = Document()

    add_section(
        document,
        "Query Decomposition and Planning",
        [
            "Complex retrieval questions may contain several information "
            "needs. Query decomposition creates smaller retrieval-oriented "
            "subqueries while retaining the original query as an important "
            "reference.",
            "Decomposition can improve recall when different parts of a "
            "question correspond to different sections or documents.",
        ],
    )

    add_section(
        document,
        "Subquery Generation",
        repeated_paragraphs(
            "Query decomposition",
            [
                "subquery generation",
                "multi-hop retrieval",
                "information needs",
                "query planning",
            ],
            220,
        ),
    )

    add_section(
        document,
        "Decomposition Risks",
        [
            "Poor decomposition can remove important constraints from the "
            "original query. It can also produce overly broad subqueries "
            "that increase the candidate pool with irrelevant material.",
            "A deterministic evaluation benchmark should record expected "
            "subqueries so that retrieval behavior can be reproduced.",
        ],
    )

    add_section(
        document,
        "Multi-hop Retrieval",
        repeated_paragraphs(
            "Multi-hop retrieval",
            [
                "cross-document evidence",
                "intermediate facts",
                "dependent evidence",
                "answer synthesis",
            ],
            190,
        ),
    )

    add_section(
        document,
        "Planning and Ranking",
        repeated_paragraphs(
            "Retrieval planning",
            [
                "original query",
                "subquery results",
                "deduplication",
                "reranking",
            ],
            170,
        ),
    )

    document.save(OUTPUT_DIR / "query_decomposition_and_planning.docx")


def create_evaluation_methodology_document() -> None:
    document = Document()

    # Deliberately headless: no Word heading styles.
    add_paragraphs(
        document,
        [
            "Retrieval Evaluation Methodology",
            "This document describes a reproducible methodology for evaluating "
            "retrieval systems against a fixed corpus and explicit relevance "
            "judgments.",
            "Evaluation should distinguish candidate generation from final "
            "ranking. A relevant chunk that never enters the candidate pool "
            "represents a candidate-generation failure.",
        ],
    )

    add_paragraphs(
        document,
        repeated_paragraphs(
            "Evaluation methodology",
            [
                "macro-averaged recall",
                "macro-averaged precision",
                "hit rate",
                "mean reciprocal rank",
            ],
            230,
        ),
    )

    add_paragraphs(
        document,
        [
            "Relevance can be represented using a graded scale. A score of "
            "three indicates direct or essential evidence, two indicates "
            "strong supporting evidence, one indicates related but "
            "insufficient evidence, and zero indicates irrelevance.",
            "The binary metrics in the initial benchmark treat scores of one "
            "or higher as relevant. The original graded labels remain "
            "available for later metrics such as NDCG.",
        ],
    )

    rows = [
        [
            f"case-{i:03d}",
            str((i % 10) + 1),
            str((i % 7) + 1),
            str(3 if i % 4 == 0 else 2),
            "candidate" if i % 3 else "reranked",
        ]
        for i in range(1, 151)
    ]

    add_table(
        document,
        ["Case", "Relevant Chunks", "Retrieved", "Relevance", "Stage"],
        rows,
    )

    add_paragraphs(
        document,
        repeated_paragraphs(
            "Benchmark interpretation",
            [
                "retrieval recall",
                "precision at k",
                "candidate coverage",
                "reranker contribution",
            ],
            200,
        ),
    )

    document.save(OUTPUT_DIR / "retrieval_evaluation_methodology.docx")


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    create_vector_search_fundamentals()
    create_hnsw_document()
    create_reranking_document()
    create_query_decomposition_document()
    create_evaluation_methodology_document()

    print(f"Generated DOCX benchmark corpus in {OUTPUT_DIR}")


if __name__ == "__main__":
    main()
