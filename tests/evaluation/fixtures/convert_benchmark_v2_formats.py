"""One-off generator: rebuilds selected Benchmark V2 source documents as
genuine .docx / .pdf files from the existing, already-verified Markdown
prose, so the benchmark exercises the real DocxLoader / PDFLoader rather
than Markdown-only ingestion.

This is NOT part of the test suite. It is run once (and re-run only if the
source content changes) to produce the committed binary fixtures under
tests/evaluation/fixtures/benchmark_v2/. Mirrors the existing
tests/evaluation/fixtures/generate_benchmark_docx.py pattern for V1.

Markdown -> DOCX: headings become native python-docx "Heading N" paragraph
styles (so DocxLoader's own `style_name.startswith("Heading ")` check
recognizes them), tables become real python-docx tables (real rows/cells,
not text), paragraphs become plain Word paragraphs, list items become
plain paragraphs prefixed with a literal "- " so StructureExtractor's list
pattern still recognizes them once DocxLoader round-trips the text.

Markdown -> PDF: built with reportlab as a genuinely flowed document
(headings as larger/bold Paragraph flowables, body text as ordinary
Paragraph flowables, tables as real Table flowables, explicit page breaks
for multi-page sections). No artificial whitespace tricks are used to
coax a particular extraction shape out of pypdf -- the resulting
structural fidelity (or lack of it) is left to reflect pypdf's actual,
unmodified `extract_text()` behavior, which is the whole point of testing
PDF ingestion honestly.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

SOURCES_DIR = Path(__file__).parent / "benchmark_v2_sources"
FIXTURES_DIR = Path(__file__).parent / "benchmark_v2"

_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
_LIST_RE = re.compile(r"^[-*+]\s+(.+)$")
_TABLE_ROW_RE = re.compile(r"^\|?(.+\|.+?)\|?$")
_TABLE_SEP_RE = re.compile(r"^\|?\s*:?-+:?\s*(\|\s*:?-+:?\s*)+\|?$")


@dataclass
class Block:
    kind: str  # "heading" | "paragraph" | "list_item" | "table"
    level: int = 0
    text: str = ""
    headers: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)


def _parse_table_row(line: str) -> list[str]:
    trimmed = line.strip()
    if trimmed.startswith("|"):
        trimmed = trimmed[1:]
    if trimmed.endswith("|"):
        trimmed = trimmed[:-1]
    return [cell.strip() for cell in trimmed.split("|")]


def parse_markdown_blocks(markdown_text: str) -> list[Block]:
    """Parse one of our own well-formed benchmark .md files into blocks.

    Deliberately simple (not a general Markdown parser): tailored to the
    exact conventions used across this corpus -- one paragraph per
    physical line, blank-line-separated, `#`-style headings, GFM pipe
    tables with a header + separator row.
    """
    lines = markdown_text.splitlines()
    blocks: list[Block] = []
    index = 0

    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()

        if not stripped:
            index += 1
            continue

        heading_match = _HEADING_RE.match(stripped)
        if heading_match:
            blocks.append(
                Block(
                    kind="heading",
                    level=len(heading_match.group(1)),
                    text=heading_match.group(2),
                )
            )
            index += 1
            continue

        table_match = (
            _TABLE_ROW_RE.match(stripped)
            and index + 1 < len(lines)
            and _TABLE_SEP_RE.match(lines[index + 1].strip())
        )
        if table_match:
            headers = _parse_table_row(stripped)
            cursor = index + 2
            rows: list[list[str]] = []
            while cursor < len(lines):
                candidate = lines[cursor].strip()
                if not candidate or not _TABLE_ROW_RE.match(candidate):
                    break
                rows.append(_parse_table_row(candidate))
                cursor += 1
            blocks.append(Block(kind="table", headers=headers, rows=rows))
            index = cursor
            continue

        list_match = _LIST_RE.match(stripped)
        if list_match:
            blocks.append(Block(kind="list_item", text=list_match.group(1)))
            index += 1
            continue

        blocks.append(Block(kind="paragraph", text=stripped))
        index += 1

    return blocks


def build_docx(blocks: list[Block], output_path: Path, title: str) -> None:
    from docx import Document as DocxDocument
    from docx.enum.table import WD_TABLE_ALIGNMENT

    document = DocxDocument()
    document.add_heading(title, level=0)  # "Title" style, not "Heading N"

    for block in blocks:
        if block.kind == "heading":
            document.add_heading(block.text, level=block.level)
        elif block.kind == "paragraph":
            document.add_paragraph(block.text)
        elif block.kind == "list_item":
            document.add_paragraph(f"- {block.text}")
        elif block.kind == "table":
            table = document.add_table(rows=1, cols=len(block.headers))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.LEFT
            for col_index, header in enumerate(block.headers):
                table.rows[0].cells[col_index].text = header
            for row in block.rows:
                cells = table.add_row().cells
                for col_index, value in enumerate(row):
                    cells[col_index].text = value

    document.save(output_path)


def build_pdf(blocks: list[Block], output_path: Path, title: str) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import (
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("DocTitle", parent=styles["Title"], fontSize=20)
    heading_styles = {
        1: ParagraphStyle("H1", parent=styles["Heading1"], fontSize=16),
        2: ParagraphStyle("H2", parent=styles["Heading2"], fontSize=13),
        3: ParagraphStyle("H3", parent=styles["Heading3"], fontSize=11),
    }
    body_style = ParagraphStyle("Body", parent=styles["Normal"], fontSize=10, leading=14)
    list_style = ParagraphStyle("ListItem", parent=body_style, leftIndent=14)

    doc = SimpleDocTemplate(
        str(output_path),
        pagesize=LETTER,
        topMargin=54,
        bottomMargin=54,
        leftMargin=56,
        rightMargin=56,
    )

    story = [Paragraph(title, title_style), Spacer(1, 10)]

    # Force a page break roughly every N top-level (level-2) headings so
    # the document is genuinely multi-page and some sections' content
    # continues across a page boundary, without hand-tuning *which*
    # sentence falls where -- reportlab's own pagination decides that.
    top_level_heading_count = 0

    for block in blocks:
        if block.kind == "heading":
            if block.level == 2:
                top_level_heading_count += 1
                if top_level_heading_count > 1 and top_level_heading_count % 3 == 1:
                    story.append(PageBreak())
            style = heading_styles.get(block.level, heading_styles[3])
            story.append(Paragraph(block.text, style))
        elif block.kind == "paragraph":
            story.append(Paragraph(block.text, body_style))
        elif block.kind == "list_item":
            story.append(Paragraph(f"- {block.text}", list_style))
        elif block.kind == "table":
            data = [block.headers, *block.rows]
            table = Table(data, hAlign="LEFT")
            table.setStyle(
                TableStyle(
                    [
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                        ("FONTSIZE", (0, 0), (-1, -1), 8),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ]
                )
            )
            story.append(table)

    doc.build(story)


CONVERSIONS: list[tuple[str, str, str, str]] = [
    # (source_md_filename, output_filename, output_format, title)
    (
        "embeddings_for_retrieval.md",
        "embeddings_for_retrieval.docx",
        "docx",
        "Embedding Models for Semantic Retrieval",
    ),
    (
        "postgres_pgvector_operations.md",
        "postgres_pgvector_operations.docx",
        "docx",
        "PostgreSQL and pgvector for Production Vector Search",
    ),
    (
        "retrieval_evaluation_metrics.md",
        "retrieval_evaluation_metrics.docx",
        "docx",
        "Evaluation Metrics and Benchmark Design for Retrieval Systems",
    ),
    (
        "chunking_strategies.md",
        "chunking_strategies.pdf",
        "pdf",
        "Chunking Strategies for Retrieval-Augmented Generation",
    ),
    (
        "rag_pipeline_architecture.md",
        "rag_pipeline_architecture.pdf",
        "pdf",
        "Architecture of a Production RAG Retrieval Pipeline",
    ),
    (
        "document_ingestion_pipeline.md",
        "document_ingestion_pipeline.pdf",
        "pdf",
        "Document Ingestion and Structural Extraction for RAG",
    ),
]


def main() -> None:
    for source_name, output_name, fmt, title in CONVERSIONS:
        source_path = SOURCES_DIR / source_name
        output_path = FIXTURES_DIR / output_name

        markdown_text = source_path.read_text(encoding="utf-8")
        # Drop the leading "# Title" block -- it becomes the native
        # Title/Heading-0 element in both DOCX and PDF instead.
        first_heading = _HEADING_RE.match(markdown_text.splitlines()[0].strip())
        body_text = (
            "\n".join(markdown_text.splitlines()[1:])
            if first_heading and first_heading.group(1) == "#"
            else markdown_text
        )

        blocks = parse_markdown_blocks(body_text)
        print(f"{source_name}: parsed {len(blocks)} blocks -> {output_name} ({fmt})")

        if fmt == "docx":
            build_docx(blocks, output_path, title)
        elif fmt == "pdf":
            build_pdf(blocks, output_path, title)
        else:
            raise ValueError(f"Unknown format: {fmt}")


if __name__ == "__main__":
    main()
