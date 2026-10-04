"""API contract tests for the document-management HTTP endpoints.

These tests pin down the external HTTP contract for the document-management
API approved in this checkpoint, *before* any routes are implemented. They
are intentionally written against endpoints that do not exist yet
(`app.main.app` currently only mounts `/research` and `/retrieval`), so they
are expected to fail today and should start passing, unmodified, once the
routes described below are implemented in a later checkpoint.

Approved routes and their explicit identifier semantics:

    GET    /documents
    GET    /documents/{logical_document_id}
    GET    /documents/{logical_document_id}/versions
    POST   /documents
    POST   /documents/{logical_document_id}/versions/{version_id}/retry
    DELETE /documents/{logical_document_id}/versions/{version_id}
    DELETE /documents/{logical_document_id}

Our domain has two distinct identifiers, and every route below names
whichever one(s) it actually operates on -- there is no overloaded or
ambiguous `{document_id}`:
  - `logical_document_id` (`Document.logical_document_id`) is the stable
    identity of a document across all of its ingested versions. Routes that
    operate on the logical document as a whole (list, get, list-versions,
    delete-the-whole-document) take only this.
  - `version_id` is one specific ingested version -- the primary key of a
    single `Document` row (what the rest of the codebase already calls a
    "document_id", e.g. `DocumentRepository.get_by_id`,
    `DocumentDeletionService.delete_version`,
    `IngestionService.retry_document`). Retry and single-version delete
    target a specific version directly via this id; they do not need to
    resolve "which version" from the logical id.

Anticipated dependency providers (to be added to `app.core.dependencies`
alongside the routes):
  - get_ingestion_service         -> IngestionService
        used by POST /documents and
        POST /documents/{logical_document_id}/versions/{version_id}/retry
  - get_document_repository       -> DocumentRepository
        used by the GET routes, and to validate that `version_id` belongs
        to `logical_document_id` before retrying or deleting a version
  - get_document_deletion_service -> DocumentDeletionService
        used by the DELETE routes

These are imported defensively below: while they don't exist, dependency
overrides are simply skipped and the real (currently 404) routing behavior
is exercised instead, so this module can be run both before and after the
routes exist without modification.

Per the approved checkpoint, these are API *contract* tests, not full
ingestion integration tests: every service dependency is faked so no
database, embedding model, or filesystem is touched.
"""

from __future__ import annotations

import io
from datetime import UTC, datetime
from uuid import UUID, uuid4

import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient

from app.core.models import Document, DocumentStatus
from app.main import app


def _optional_dependency(name: str):
    """Import one anticipated dependency provider if it exists yet.

    Each of the three dependency providers lands in its own checkpoint, so
    these are imported independently rather than in one statement -- one
    provider existing must not be masked by another one not existing yet.
    """

    try:
        module = __import__("app.core.dependencies", fromlist=[name])
        return getattr(module, name)
    except (ImportError, AttributeError):
        return None


get_ingestion_service = _optional_dependency("get_ingestion_service")
get_document_repository = _optional_dependency("get_document_repository")
get_document_deletion_service = _optional_dependency("get_document_deletion_service")


# ---------------------------------------------------------------------------
# Test doubles
# ---------------------------------------------------------------------------


def make_document(
    *,
    document_id: UUID | None = None,
    logical_document_id: UUID | None = None,
    version_number: int = 1,
    is_current: bool = True,
    status: DocumentStatus = DocumentStatus.READY,
    title: str = "Sample Document",
    document_type: str = "general",
    source: str | None = None,
    processing_attempt: int = 1,
    last_error: str | None = None,
    failed_at: datetime | None = None,
    created_at: datetime | None = None,
) -> Document:
    """Build an in-memory `Document` row for use in fakes (never persisted).

    Reuses the real ORM class purely as a typed data container so the fakes
    can never drift from the real column/attribute names.
    """

    now = datetime.now(UTC)
    created = created_at or now

    return Document(
        id=document_id or uuid4(),
        logical_document_id=logical_document_id or uuid4(),
        title=title,
        document_type=document_type,
        source=source,
        content_hash=f"hash-{uuid4()}",
        version_number=version_number,
        is_current=is_current,
        status=status,
        processing_attempt=processing_attempt,
        processing_started_at=now,
        processing_completed_at=now if status == DocumentStatus.READY else None,
        failed_at=failed_at,
        last_error=last_error,
        document_metadata={},
        created_at=created,
        updated_at=created,
    )


class FakeIngestionService:
    """Stands in for `app.ingestion.service.IngestionService`.

    Mirrors its two real public methods exactly (`ingest_file`,
    `retry_document`) so whichever route implementation lands later is
    exercised through the same call shape it will use in production.
    """

    def __init__(self) -> None:
        self.ingest_calls: list[dict[str, object]] = []
        self.retry_calls: list[UUID] = []
        self.next_document: Document | None = None
        self.retry_result: Document | None = None
        self.raise_on_retry: Exception | None = None
        self.raise_on_ingest: Exception | None = None

    async def ingest_file(
        self,
        path,
        loader,
        title,
        document_type,
        logical_document_id=None,
        source=None,
        original_filename=None,
    ) -> Document:
        self.ingest_calls.append(
            {
                "title": title,
                "document_type": document_type,
                "logical_document_id": logical_document_id,
                "source": source,
                "original_filename": original_filename,
            }
        )
        if self.raise_on_ingest is not None:
            # Mirrors the real IngestionService.ingest_file contract: on a
            # processing failure it persists a FAILED Document and then
            # raises, with the persisted row attached to the exception as
            # `.failed_document` so the route can recover and report it.
            raise self.raise_on_ingest
        assert self.next_document is not None, "Fake not configured with next_document"
        return self.next_document

    async def retry_document(self, document_id: UUID, loader=None) -> Document:
        self.retry_calls.append(document_id)
        if self.raise_on_retry is not None:
            raise self.raise_on_retry
        assert self.retry_result is not None, "Fake not configured with retry_result"
        return self.retry_result


class FakeDocumentRepository:
    """Stands in for `app.core.repositories.document.DocumentRepository`.

    `get_by_id`, `get_by_logical_document_id`, `get_current_version`, and
    `list_logical_documents` all mirror the real repository's methods.
    """

    def __init__(self, documents: list[Document] | None = None) -> None:
        self.documents = list(documents or [])

    async def list_logical_documents(self) -> list[tuple[UUID, Document | None]]:
        logical_document_ids = sorted(
            {doc.logical_document_id for doc in self.documents},
            key=str,
        )
        return [
            (
                logical_document_id,
                next(
                    (
                        d
                        for d in self.documents
                        if d.logical_document_id == logical_document_id and d.is_current
                    ),
                    None,
                ),
            )
            for logical_document_id in logical_document_ids
        ]

    async def get_by_id(self, document_id: UUID) -> Document | None:
        return next((d for d in self.documents if d.id == document_id), None)

    async def get_by_logical_document_id(
        self,
        logical_document_id: UUID,
    ) -> list[Document]:
        versions = [d for d in self.documents if d.logical_document_id == logical_document_id]
        return sorted(versions, key=lambda d: d.version_number, reverse=True)

    async def get_current_version(self, logical_document_id: UUID) -> Document | None:
        return next(
            (
                d
                for d in self.documents
                if d.logical_document_id == logical_document_id and d.is_current
            ),
            None,
        )


class FakeDocumentDeletionService:
    """Stands in for `app.core.services.document_deletion.DocumentDeletionService`.

    Mirrors its two real public methods exactly.
    """

    def __init__(self) -> None:
        self.deleted_versions: list[UUID] = []
        self.deleted_logical_documents: list[UUID] = []
        self.raise_on_delete_version: Exception | None = None
        self.raise_on_delete_logical: Exception | None = None

    async def delete_version(self, document_id: UUID) -> None:
        if self.raise_on_delete_version is not None:
            raise self.raise_on_delete_version
        self.deleted_versions.append(document_id)

    async def delete_logical_document(self, logical_document_id: UUID) -> None:
        if self.raise_on_delete_logical is not None:
            raise self.raise_on_delete_logical
        self.deleted_logical_documents.append(logical_document_id)


def override(dependency, fake) -> None:
    """Register a dependency override, a no-op while `dependency` is None."""

    if dependency is not None:
        app.dependency_overrides[dependency] = lambda: fake


@pytest.fixture(autouse=True)
def _clear_dependency_overrides():
    yield
    app.dependency_overrides.clear()


@pytest.fixture
def client() -> TestClient:
    return TestClient(app)


# ---------------------------------------------------------------------------
# Response-shape assertions
# ---------------------------------------------------------------------------

EXPECTED_VERSION_FIELDS = {
    "id",
    "logical_document_id",
    "version_number",
    "is_current",
    "status",
    "title",
    "document_type",
    "source",
    "processing_attempt",
    "created_at",
    "updated_at",
    "processing_started_at",
    "processing_completed_at",
    "failed_at",
    "last_error",
}

# Fields that exist on the SQLAlchemy model but must never leak into the
# HTTP response.
FORBIDDEN_FIELDS = {
    "_sa_instance_state",
    "metadata",
    "document_metadata",
    "content_hash",
    "authors",
    "publication_date",
    "pages",
    "chunks",
    "sections",
}


def assert_version_response_shape(body: dict) -> None:
    assert EXPECTED_VERSION_FIELDS.issubset(body.keys())
    assert FORBIDDEN_FIELDS.isdisjoint(body.keys())


# GET /documents represents LOGICAL documents, not individual versions: each
# item is a logical_document_id plus its (nullable) current version. The
# nested current_version reuses the exact same shape as DocumentVersionResponse
# (field `id`, same as POST /documents) -- there is only one version-shaped
# response contract, not a near-duplicate with a differently named identifier.
EXPECTED_LOGICAL_DOCUMENT_FIELDS = {"logical_document_id", "current_version"}


def assert_logical_document_response_shape(item: dict) -> None:
    assert set(item.keys()) == EXPECTED_LOGICAL_DOCUMENT_FIELDS

    if item["current_version"] is not None:
        assert_version_response_shape(item["current_version"])


def upload_file(client: TestClient, **form_overrides):
    data = {"document_type": "research_paper"}
    data.update(form_overrides)
    files = {"file": ("paper.md", b"# Title\n\nSome ingested content.", "text/markdown")}
    return client.post("/documents", data=data, files=files)


def build_docx_bytes(paragraphs: list[str]) -> bytes:
    """Build real, parseable .docx bytes in-memory via python-docx -- the
    same "build a tiny real fixture with the format's own library"
    convention already used for PDF loader tests, just assembled in memory
    rather than written to a temp path since this is only ever POSTed."""

    document = DocxDocument()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# 1. Successful document upload
# ---------------------------------------------------------------------------


def test_upload_document_returns_created_version(client: TestClient) -> None:
    ingestion = FakeIngestionService()
    document = make_document(
        title="My Paper",
        document_type="research_paper",
        status=DocumentStatus.PROCESSING,
        is_current=False,
        version_number=1,
    )
    ingestion.next_document = document

    override(get_ingestion_service, ingestion)

    response = upload_file(client, title="My Paper")

    assert response.status_code == 201

    body = response.json()
    assert_version_response_shape(body)
    assert body["id"] == str(document.id)
    assert body["logical_document_id"] == str(document.logical_document_id)
    assert body["title"] == "My Paper"
    assert body["document_type"] == "research_paper"
    assert body["version_number"] == 1
    assert body["status"] in {"PROCESSING", "READY"}

    assert len(ingestion.ingest_calls) == 1
    assert ingestion.ingest_calls[0]["title"] == "My Paper"
    assert ingestion.ingest_calls[0]["document_type"] == "research_paper"

    # The client's actual uploaded filename must reach IngestionService,
    # not whatever temporary filesystem path the route used internally.
    assert ingestion.ingest_calls[0]["original_filename"] == "paper.md"


def test_upload_accepts_docx_file_and_reaches_ready(client: TestClient) -> None:
    """POST /documents accepts a real .docx upload, passes the client's
    actual original filename through to IngestionService (not the route's
    internal temp path), and reports successful ingestion as READY using
    the existing fake-ingestion-service infrastructure -- the route itself
    never touches the real ingestion/embedding pipeline in this test."""

    ingestion = FakeIngestionService()
    document = make_document(
        title="Word Paper",
        document_type="research_paper",
        status=DocumentStatus.READY,
        is_current=True,
    )
    ingestion.next_document = document

    override(get_ingestion_service, ingestion)

    docx_bytes = build_docx_bytes(
        [
            "Retrieval-Augmented Generation",
            "RAG combines retrieval with generation.",
        ]
    )
    files = {
        "file": (
            "paper.docx",
            docx_bytes,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    response = client.post(
        "/documents",
        data={"document_type": "research_paper", "title": "Word Paper"},
        files=files,
    )

    assert response.status_code == 201

    body = response.json()
    assert_version_response_shape(body)
    assert body["status"] == "READY"

    assert len(ingestion.ingest_calls) == 1
    assert ingestion.ingest_calls[0]["original_filename"] == "paper.docx"


# ---------------------------------------------------------------------------
# 2. Duplicate / idempotent upload
# ---------------------------------------------------------------------------


def test_duplicate_upload_returns_same_version(client: TestClient) -> None:
    ingestion = FakeIngestionService()
    document = make_document(title="Repeated Paper", status=DocumentStatus.READY)
    ingestion.next_document = document

    override(get_ingestion_service, ingestion)

    first_response = upload_file(client, title="Repeated Paper")
    second_response = upload_file(client, title="Repeated Paper")

    assert first_response.status_code == 201
    assert second_response.status_code == 201

    first_body = first_response.json()
    second_body = second_response.json()

    # Content-hash idempotency: re-ingesting identical content returns the
    # same underlying version rather than creating a new one.
    assert first_body["id"] == second_body["id"] == str(document.id)
    assert len(ingestion.ingest_calls) == 2


# ---------------------------------------------------------------------------
# 3. Upload with an explicitly supplied logical_document_id
# ---------------------------------------------------------------------------


def test_upload_with_explicit_logical_document_id(client: TestClient) -> None:
    ingestion = FakeIngestionService()
    logical_document_id = uuid4()
    document = make_document(
        logical_document_id=logical_document_id,
        version_number=2,
        title="New Version",
    )
    ingestion.next_document = document

    override(get_ingestion_service, ingestion)

    response = upload_file(
        client,
        title="New Version",
        logical_document_id=str(logical_document_id),
    )

    assert response.status_code == 201

    body = response.json()
    assert body["logical_document_id"] == str(logical_document_id)
    assert body["version_number"] == 2

    assert ingestion.ingest_calls[0]["logical_document_id"] == logical_document_id


# ---------------------------------------------------------------------------
# 4. Invalid request / file handling
# ---------------------------------------------------------------------------


def test_upload_without_file_is_rejected(client: TestClient) -> None:
    ingestion = FakeIngestionService()
    override(get_ingestion_service, ingestion)

    response = client.post("/documents", data={"document_type": "research_paper"})

    assert response.status_code == 422
    assert ingestion.ingest_calls == []


def test_upload_without_document_type_is_rejected(client: TestClient) -> None:
    ingestion = FakeIngestionService()
    override(get_ingestion_service, ingestion)

    files = {"file": ("paper.md", b"content", "text/markdown")}
    response = client.post("/documents", data={}, files=files)

    assert response.status_code == 422
    assert ingestion.ingest_calls == []


def test_upload_with_empty_file_is_rejected(client: TestClient) -> None:
    ingestion = FakeIngestionService()
    override(get_ingestion_service, ingestion)

    files = {"file": ("empty.md", b"", "text/markdown")}
    response = client.post("/documents", data={"document_type": "research_paper"}, files=files)

    assert response.status_code == 422
    assert ingestion.ingest_calls == []


def test_upload_with_invalid_logical_document_id_is_rejected(client: TestClient) -> None:
    ingestion = FakeIngestionService()
    override(get_ingestion_service, ingestion)

    response = upload_file(client, logical_document_id="not-a-uuid")

    assert response.status_code == 422
    assert ingestion.ingest_calls == []


def test_upload_with_unsupported_extension_is_rejected(client: TestClient) -> None:
    """Loader resolution happens in the route before IngestionService is
    ever called, so an unsupported extension is rejected the same way
    regardless of which (fake or real) ingestion service is configured."""

    ingestion = FakeIngestionService()
    override(get_ingestion_service, ingestion)

    files = {"file": ("paper.exe", b"not a real document", "application/octet-stream")}
    response = client.post("/documents", data={"document_type": "research_paper"}, files=files)

    assert response.status_code == 422
    assert ingestion.ingest_calls == []


def test_upload_with_legacy_doc_extension_is_rejected(client: TestClient) -> None:
    """Legacy .doc (pre-2007 binary Word format) is intentionally
    unsupported -- only modern .docx is accepted."""

    ingestion = FakeIngestionService()
    override(get_ingestion_service, ingestion)

    files = {"file": ("paper.doc", b"not a real document", "application/msword")}
    response = client.post("/documents", data={"document_type": "research_paper"}, files=files)

    assert response.status_code == 422
    assert ingestion.ingest_calls == []


# ---------------------------------------------------------------------------
# 5. Ingestion failure producing a FAILED document
# ---------------------------------------------------------------------------


def test_upload_that_fails_ingestion_returns_failed_document(client: TestClient) -> None:
    """IngestionService marks the version FAILED (with last_error/failed_at)
    before surfacing the failure; the route is expected to report that
    FAILED version back to the caller rather than a bare 5xx, since it is a
    normal, trackable, retry-able lifecycle outcome."""

    ingestion = FakeIngestionService()
    failed_at = datetime.now(UTC)
    document = make_document(
        status=DocumentStatus.FAILED,
        is_current=False,
        last_error="Embedding generation timed out.",
        failed_at=failed_at,
        processing_attempt=1,
    )
    ingestion.next_document = document

    override(get_ingestion_service, ingestion)

    response = upload_file(client)

    assert response.status_code == 201

    body = response.json()
    assert_version_response_shape(body)
    assert body["status"] == "FAILED"
    assert body["last_error"] == "Embedding generation timed out."
    assert body["failed_at"] is not None
    assert body["is_current"] is False


def test_upload_that_raises_during_ingestion_still_returns_201(client: TestClient) -> None:
    """Models the REAL IngestionService.ingest_file contract precisely: on a
    processing failure it persists a FAILED Document (last_error/failed_at
    set, committed) and then raises the original exception -- it does not
    return a document directly. The route has no way to re-derive which
    document that was from the exception alone unless IngestionService
    attaches it, so the service is expected to set `.failed_document` on
    the exception it raises before re-raising. The route must recover that
    attached document and still report it with 201, not a 422/500."""

    ingestion = FakeIngestionService()

    failed_document = make_document(
        status=DocumentStatus.FAILED,
        is_current=False,
        last_error="Connection reset while embedding.",
        failed_at=datetime.now(UTC),
        processing_attempt=1,
    )

    processing_error = RuntimeError("Connection reset while embedding.")
    processing_error.failed_document = failed_document
    ingestion.raise_on_ingest = processing_error

    override(get_ingestion_service, ingestion)

    response = upload_file(client)

    assert response.status_code == 201

    body = response.json()
    assert_version_response_shape(body)
    assert body["id"] == str(failed_document.id)
    assert body["status"] == "FAILED"
    assert body["last_error"] == "Connection reset while embedding."
    assert body["failed_at"] is not None
    assert body["is_current"] is False


def test_upload_that_raises_without_recoverable_document_is_not_fabricated(
    client: TestClient,
) -> None:
    """If IngestionService raises without ever persisting anything (e.g. a
    failure before any Document row existed), there is no FAILED version to
    report -- the route must not fabricate a success response, and must not
    mislabel it as a request-validation failure (422) either. It is
    acceptable for the route to simply not handle this case, letting it
    surface as a generic server error -- which is exactly what happens here
    (TestClient re-raises an exception a route doesn't handle instead of
    turning it into a response, since there is nothing meaningful for the
    route to construct a DocumentVersionResponse from)."""

    ingestion = FakeIngestionService()
    ingestion.raise_on_ingest = RuntimeError("Could not read file at all.")

    override(get_ingestion_service, ingestion)

    with pytest.raises(RuntimeError, match="Could not read file at all."):
        upload_file(client)


# ---------------------------------------------------------------------------
# 6. Document listing
# ---------------------------------------------------------------------------


def test_list_documents_returns_one_item_per_logical_document(client: TestClient) -> None:
    """Multiple versions of the SAME logical document must collapse into a
    single list item representing its current version."""

    logical_document_id = uuid4()
    old_version = make_document(
        logical_document_id=logical_document_id,
        version_number=1,
        is_current=False,
        status=DocumentStatus.READY,
        title="Doc One (old)",
    )
    current_version = make_document(
        logical_document_id=logical_document_id,
        version_number=2,
        is_current=True,
        status=DocumentStatus.READY,
        title="Doc One",
    )

    repository = FakeDocumentRepository([old_version, current_version])
    override(get_document_repository, repository)

    response = client.get("/documents")

    assert response.status_code == 200

    body = response.json()
    assert isinstance(body, list)
    assert len(body) == 1

    item = body[0]
    assert_logical_document_response_shape(item)
    assert item["logical_document_id"] == str(logical_document_id)

    # The current READY version is represented correctly, using the same
    # canonical `id` field as POST /documents (not a differently named
    # `version_id`).
    assert item["current_version"]["id"] == str(current_version.id)
    assert item["current_version"]["logical_document_id"] == str(logical_document_id)
    assert item["current_version"]["version_number"] == 2
    assert item["current_version"]["status"] == "READY"
    assert item["current_version"]["is_current"] is True
    assert item["current_version"]["title"] == "Doc One"


def test_list_documents_returns_all_logical_documents(client: TestClient) -> None:
    """Multiple distinct logical documents are all returned, with no
    duplicate logical_document_id values."""

    doc_one = make_document(title="Doc One", is_current=True)
    doc_two = make_document(title="Doc Two", is_current=True)
    doc_three = make_document(title="Doc Three", is_current=True)

    repository = FakeDocumentRepository([doc_one, doc_two, doc_three])
    override(get_document_repository, repository)

    response = client.get("/documents")

    assert response.status_code == 200

    body = response.json()
    assert len(body) == 3

    returned_logical_ids = [item["logical_document_id"] for item in body]
    assert len(returned_logical_ids) == len(set(returned_logical_ids))
    assert set(returned_logical_ids) == {
        str(doc_one.logical_document_id),
        str(doc_two.logical_document_id),
        str(doc_three.logical_document_id),
    }


def test_list_documents_represents_logical_document_with_no_current_version(
    client: TestClient,
) -> None:
    """A logical document whose only version(s) never became current (e.g.
    still PROCESSING, or FAILED) must still appear in the list, with
    current_version=None -- it is not silently omitted."""

    logical_document_id = uuid4()
    failed_version = make_document(
        logical_document_id=logical_document_id,
        status=DocumentStatus.FAILED,
        is_current=False,
        last_error="Embedding provider timed out.",
    )

    repository = FakeDocumentRepository([failed_version])
    override(get_document_repository, repository)

    response = client.get("/documents")

    assert response.status_code == 200

    body = response.json()
    assert len(body) == 1

    item = body[0]
    assert_logical_document_response_shape(item)
    assert item["logical_document_id"] == str(logical_document_id)
    assert item["current_version"] is None


def test_list_documents_failed_and_processing_versions_are_never_current(
    client: TestClient,
) -> None:
    """A logical document with a READY current version plus a newer FAILED
    attempt must still report the READY one as current -- FAILED/PROCESSING
    versions must never be mistaken for the current version."""

    logical_document_id = uuid4()
    ready_current = make_document(
        logical_document_id=logical_document_id,
        version_number=1,
        is_current=True,
        status=DocumentStatus.READY,
    )
    failed_attempt = make_document(
        logical_document_id=logical_document_id,
        version_number=2,
        is_current=False,
        status=DocumentStatus.FAILED,
    )
    processing_attempt = make_document(
        logical_document_id=logical_document_id,
        version_number=3,
        is_current=False,
        status=DocumentStatus.PROCESSING,
    )

    repository = FakeDocumentRepository([ready_current, failed_attempt, processing_attempt])
    override(get_document_repository, repository)

    response = client.get("/documents")

    assert response.status_code == 200

    body = response.json()
    assert len(body) == 1

    current_version = body[0]["current_version"]
    assert current_version["id"] == str(ready_current.id)
    assert current_version["status"] == "READY"
    assert current_version["is_current"] is True


def test_list_documents_empty_library(client: TestClient) -> None:
    override(get_document_repository, FakeDocumentRepository([]))

    response = client.get("/documents")

    assert response.status_code == 200
    assert response.json() == []


def test_list_documents_current_version_uses_canonical_version_schema(
    client: TestClient,
) -> None:
    """current_version must be exactly the same shape as DocumentVersionResponse
    (field `id`, same as POST /documents) -- not a near-duplicate schema with
    a differently named identifier, and nothing SQLAlchemy-internal leaks
    through."""

    document = make_document()

    override(get_document_repository, FakeDocumentRepository([document]))

    response = client.get("/documents")

    assert response.status_code == 200

    item = response.json()[0]
    assert set(item.keys()) == EXPECTED_LOGICAL_DOCUMENT_FIELDS
    assert set(item["current_version"].keys()) == EXPECTED_VERSION_FIELDS
    assert item["current_version"]["id"] == str(document.id)


# ---------------------------------------------------------------------------
# 7. Retrieving one document
# ---------------------------------------------------------------------------


def test_get_document_returns_current_version(client: TestClient) -> None:
    """Same response shape as one item from GET /documents: a
    logical_document_id plus its nested canonical current_version."""

    logical_document_id = uuid4()
    document = make_document(
        logical_document_id=logical_document_id,
        is_current=True,
        status=DocumentStatus.READY,
    )

    override(get_document_repository, FakeDocumentRepository([document]))

    response = client.get(f"/documents/{logical_document_id}")

    assert response.status_code == 200

    body = response.json()
    assert_logical_document_response_shape(body)
    assert body["logical_document_id"] == str(logical_document_id)
    assert body["current_version"]["id"] == str(document.id)
    assert body["current_version"]["status"] == "READY"
    assert body["current_version"]["is_current"] is True


def test_get_document_with_no_current_version_returns_null(client: TestClient) -> None:
    """An existing logical document whose only version never became current
    (e.g. still FAILED) returns 200 with current_version=None, rather than a
    404 -- the logical document itself does exist."""

    logical_document_id = uuid4()
    failed_version = make_document(
        logical_document_id=logical_document_id,
        status=DocumentStatus.FAILED,
        is_current=False,
    )

    override(get_document_repository, FakeDocumentRepository([failed_version]))

    response = client.get(f"/documents/{logical_document_id}")

    assert response.status_code == 200

    body = response.json()
    assert_logical_document_response_shape(body)
    assert body["logical_document_id"] == str(logical_document_id)
    assert body["current_version"] is None


def test_get_document_unknown_logical_document_id_returns_404(client: TestClient) -> None:
    override(get_document_repository, FakeDocumentRepository([]))

    response = client.get(f"/documents/{uuid4()}")

    assert response.status_code == 404


def test_get_document_failed_and_processing_versions_are_never_current(
    client: TestClient,
) -> None:
    """A READY current version plus newer FAILED/PROCESSING attempts must
    still report the READY one as current."""

    logical_document_id = uuid4()
    ready_current = make_document(
        logical_document_id=logical_document_id,
        version_number=1,
        is_current=True,
        status=DocumentStatus.READY,
    )
    failed_attempt = make_document(
        logical_document_id=logical_document_id,
        version_number=2,
        is_current=False,
        status=DocumentStatus.FAILED,
    )
    processing_attempt = make_document(
        logical_document_id=logical_document_id,
        version_number=3,
        is_current=False,
        status=DocumentStatus.PROCESSING,
    )

    override(
        get_document_repository,
        FakeDocumentRepository([ready_current, failed_attempt, processing_attempt]),
    )

    response = client.get(f"/documents/{logical_document_id}")

    assert response.status_code == 200

    current_version = response.json()["current_version"]
    assert current_version["id"] == str(ready_current.id)
    assert current_version["status"] == "READY"
    assert current_version["is_current"] is True


def test_get_document_response_shape_matches_list_item_exactly(client: TestClient) -> None:
    """Exact response shape -- identical to one item from GET /documents --
    and no SQLAlchemy internals leak through."""

    logical_document_id = uuid4()
    document = make_document(logical_document_id=logical_document_id)

    override(get_document_repository, FakeDocumentRepository([document]))

    response = client.get(f"/documents/{logical_document_id}")

    assert response.status_code == 200

    body = response.json()
    assert set(body.keys()) == EXPECTED_LOGICAL_DOCUMENT_FIELDS
    assert set(body["current_version"].keys()) == EXPECTED_VERSION_FIELDS
    assert FORBIDDEN_FIELDS.isdisjoint(body["current_version"].keys())


# ---------------------------------------------------------------------------
# 8. Retrieving document versions
# ---------------------------------------------------------------------------


def test_get_document_versions_lists_all_versions(client: TestClient) -> None:
    logical_document_id = uuid4()
    version_one = make_document(
        logical_document_id=logical_document_id,
        version_number=1,
        is_current=False,
        status=DocumentStatus.READY,
    )
    version_two = make_document(
        logical_document_id=logical_document_id,
        version_number=2,
        is_current=True,
        status=DocumentStatus.READY,
    )

    override(get_document_repository, FakeDocumentRepository([version_one, version_two]))

    response = client.get(f"/documents/{logical_document_id}/versions")

    assert response.status_code == 200

    body = response.json()
    assert isinstance(body, list)
    assert len(body) == 2

    for item in body:
        assert_version_response_shape(item)

    version_numbers = {item["version_number"] for item in body}
    assert version_numbers == {1, 2}


def test_get_document_versions_returns_newest_first(client: TestClient) -> None:
    """Results must be ordered by version_number DESC."""

    logical_document_id = uuid4()
    version_one = make_document(logical_document_id=logical_document_id, version_number=1)
    version_two = make_document(logical_document_id=logical_document_id, version_number=2)
    version_three = make_document(logical_document_id=logical_document_id, version_number=3)

    # Inserted out of version-number order to prove the response is sorted,
    # not merely returned in insertion order.
    repository = FakeDocumentRepository([version_two, version_three, version_one])
    override(get_document_repository, repository)

    response = client.get(f"/documents/{logical_document_id}/versions")

    assert response.status_code == 200

    body = response.json()
    assert [item["version_number"] for item in body] == [3, 2, 1]


def test_get_document_versions_ordering_is_by_version_number_not_timestamp(
    client: TestClient,
) -> None:
    """Ordering must key on version_number, not created_at/updated_at --
    even when timestamps disagree with version_number order."""

    logical_document_id = uuid4()

    # version_number 1 was (unrealistically) created AFTER version_number 2,
    # to prove sorting does not fall back to timestamp order.
    earlier_timestamp = datetime(2026, 1, 1, tzinfo=UTC)
    later_timestamp = datetime(2026, 6, 1, tzinfo=UTC)

    version_two = make_document(
        logical_document_id=logical_document_id,
        version_number=2,
        created_at=earlier_timestamp,
    )
    version_one = make_document(
        logical_document_id=logical_document_id,
        version_number=1,
        created_at=later_timestamp,
    )

    repository = FakeDocumentRepository([version_one, version_two])
    override(get_document_repository, repository)

    response = client.get(f"/documents/{logical_document_id}/versions")

    assert response.status_code == 200

    body = response.json()
    assert [item["version_number"] for item in body] == [2, 1]


def test_get_document_versions_includes_every_lifecycle_status(client: TestClient) -> None:
    """All statuses are returned -- no filtering to current or READY-only."""

    logical_document_id = uuid4()
    ready_version = make_document(
        logical_document_id=logical_document_id,
        version_number=1,
        is_current=True,
        status=DocumentStatus.READY,
    )
    failed_version = make_document(
        logical_document_id=logical_document_id,
        version_number=2,
        is_current=False,
        status=DocumentStatus.FAILED,
        last_error="Embedding timed out.",
    )
    processing_version = make_document(
        logical_document_id=logical_document_id,
        version_number=3,
        is_current=False,
        status=DocumentStatus.PROCESSING,
    )
    uploaded_version = make_document(
        logical_document_id=logical_document_id,
        version_number=4,
        is_current=False,
        status=DocumentStatus.UPLOADED,
    )

    repository = FakeDocumentRepository(
        [ready_version, failed_version, processing_version, uploaded_version]
    )
    override(get_document_repository, repository)

    response = client.get(f"/documents/{logical_document_id}/versions")

    assert response.status_code == 200

    body = response.json()
    assert len(body) == 4

    statuses_by_version_number = {item["version_number"]: item["status"] for item in body}
    assert statuses_by_version_number == {
        1: "READY",
        2: "FAILED",
        3: "PROCESSING",
        4: "UPLOADED",
    }


def test_get_document_versions_preserves_is_current_flag(client: TestClient) -> None:
    logical_document_id = uuid4()
    current_version = make_document(
        logical_document_id=logical_document_id,
        version_number=2,
        is_current=True,
        status=DocumentStatus.READY,
    )
    old_version = make_document(
        logical_document_id=logical_document_id,
        version_number=1,
        is_current=False,
        status=DocumentStatus.READY,
    )

    repository = FakeDocumentRepository([current_version, old_version])
    override(get_document_repository, repository)

    response = client.get(f"/documents/{logical_document_id}/versions")

    assert response.status_code == 200

    body = response.json()
    is_current_by_version_number = {item["version_number"]: item["is_current"] for item in body}
    assert is_current_by_version_number == {2: True, 1: False}


def test_get_document_versions_single_version(client: TestClient) -> None:
    logical_document_id = uuid4()
    only_version = make_document(logical_document_id=logical_document_id, version_number=1)

    repository = FakeDocumentRepository([only_version])
    override(get_document_repository, repository)

    response = client.get(f"/documents/{logical_document_id}/versions")

    assert response.status_code == 200

    body = response.json()
    assert len(body) == 1
    assert body[0]["id"] == str(only_version.id)
    assert body[0]["version_number"] == 1


def test_get_document_versions_unknown_logical_document_id_returns_404(
    client: TestClient,
) -> None:
    override(get_document_repository, FakeDocumentRepository([]))

    response = client.get(f"/documents/{uuid4()}/versions")

    assert response.status_code == 404


def test_get_document_versions_response_shape_is_canonical(client: TestClient) -> None:
    """Each item is exactly the canonical DocumentVersionResponse shape --
    the same contract as POST /documents and GET /documents's
    current_version -- not a new/different version schema, and no
    SQLAlchemy internals leak through."""

    logical_document_id = uuid4()
    document = make_document(logical_document_id=logical_document_id)

    repository = FakeDocumentRepository([document])
    override(get_document_repository, repository)

    response = client.get(f"/documents/{logical_document_id}/versions")

    assert response.status_code == 200

    body = response.json()
    assert len(body) == 1
    assert set(body[0].keys()) == EXPECTED_VERSION_FIELDS
    assert FORBIDDEN_FIELDS.isdisjoint(body[0].keys())


# ---------------------------------------------------------------------------
# 9. Retrying a FAILED document
# ---------------------------------------------------------------------------


def test_retry_failed_document_succeeds(client: TestClient) -> None:
    """Retry targets a specific version_id directly -- no resolution of
    "which version is FAILED" is needed; the caller already knows, because
    it is the same version_id it saw as FAILED in a prior list/get call."""

    logical_document_id = uuid4()
    failed_version = make_document(
        logical_document_id=logical_document_id,
        version_number=2,
        status=DocumentStatus.FAILED,
        is_current=False,
        last_error="Connection reset.",
    )

    ready_version = make_document(
        document_id=failed_version.id,
        logical_document_id=logical_document_id,
        version_number=2,
        status=DocumentStatus.READY,
        is_current=True,
        processing_attempt=2,
    )

    ingestion = FakeIngestionService()
    ingestion.retry_result = ready_version

    override(get_document_repository, FakeDocumentRepository([failed_version]))
    override(get_ingestion_service, ingestion)

    response = client.post(f"/documents/{logical_document_id}/versions/{failed_version.id}/retry")

    assert response.status_code == 200

    body = response.json()
    assert_version_response_shape(body)
    assert body["status"] == "READY"
    assert body["is_current"] is True
    assert body["processing_attempt"] == 2

    # The route passes version_id straight through to
    # IngestionService.retry_document -- no translation required.
    assert ingestion.retry_calls == [failed_version.id]


# ---------------------------------------------------------------------------
# 10. Rejecting retry of a non-FAILED document
# ---------------------------------------------------------------------------


def test_retry_rejects_non_failed_document(client: TestClient) -> None:
    """IngestionService.retry_document itself raises ValueError when the
    targeted version is not FAILED; the route maps that to 409 Conflict.
    The call still reaches the service (existence was already confirmed via
    the repository), it just fails the service's own status guard."""

    logical_document_id = uuid4()
    ready_version = make_document(
        logical_document_id=logical_document_id,
        status=DocumentStatus.READY,
        is_current=True,
    )

    ingestion = FakeIngestionService()
    ingestion.raise_on_retry = ValueError("Only FAILED documents can be retried.")

    override(get_document_repository, FakeDocumentRepository([ready_version]))
    override(get_ingestion_service, ingestion)

    response = client.post(f"/documents/{logical_document_id}/versions/{ready_version.id}/retry")

    assert response.status_code == 409
    assert ingestion.retry_calls == [ready_version.id]


def test_retry_rejects_processing_version(client: TestClient) -> None:
    """A PROCESSING version is not FAILED either -- same 409 as READY."""

    logical_document_id = uuid4()
    processing_version = make_document(
        logical_document_id=logical_document_id,
        status=DocumentStatus.PROCESSING,
        is_current=False,
    )

    ingestion = FakeIngestionService()
    ingestion.raise_on_retry = ValueError("Only FAILED documents can be retried.")

    override(get_document_repository, FakeDocumentRepository([processing_version]))
    override(get_ingestion_service, ingestion)

    response = client.post(
        f"/documents/{logical_document_id}/versions/{processing_version.id}/retry"
    )

    assert response.status_code == 409
    assert ingestion.retry_calls == [processing_version.id]


def test_retry_processing_failure_returns_200_with_failed_status(client: TestClient) -> None:
    """A retry that fails again during processing is reported like any other
    processing failure: 200, same version id, status FAILED, last_error
    populated -- not a 409/422/500. IngestionService persists the FAILED
    state and attaches it to the exception as `.failed_document`, exactly
    as it does for a first-time ingestion failure."""

    logical_document_id = uuid4()
    failed_version = make_document(
        logical_document_id=logical_document_id,
        status=DocumentStatus.FAILED,
        is_current=False,
        last_error="Connection reset.",
        processing_attempt=1,
    )

    failed_again = make_document(
        document_id=failed_version.id,
        logical_document_id=logical_document_id,
        status=DocumentStatus.FAILED,
        is_current=False,
        last_error="Connection reset again.",
        processing_attempt=2,
    )

    retry_error = RuntimeError("Connection reset again.")
    retry_error.failed_document = failed_again

    ingestion = FakeIngestionService()
    ingestion.raise_on_retry = retry_error

    override(get_document_repository, FakeDocumentRepository([failed_version]))
    override(get_ingestion_service, ingestion)

    response = client.post(f"/documents/{logical_document_id}/versions/{failed_version.id}/retry")

    assert response.status_code == 200

    body = response.json()
    assert_version_response_shape(body)
    assert body["id"] == str(failed_version.id)
    assert body["status"] == "FAILED"
    assert body["last_error"] == "Connection reset again."
    assert body["processing_attempt"] == 2

    assert ingestion.retry_calls == [failed_version.id]


def test_retry_unknown_version_for_existing_logical_document_returns_404(
    client: TestClient,
) -> None:
    """The logical document is real (has another version), but the specific
    version_id requested does not exist at all."""

    logical_document_id = uuid4()
    other_version = make_document(
        logical_document_id=logical_document_id,
        version_number=1,
        status=DocumentStatus.READY,
    )

    ingestion = FakeIngestionService()

    override(get_document_repository, FakeDocumentRepository([other_version]))
    override(get_ingestion_service, ingestion)

    response = client.post(f"/documents/{logical_document_id}/versions/{uuid4()}/retry")

    assert response.status_code == 404
    assert ingestion.retry_calls == []


def test_retry_response_shape_is_canonical(client: TestClient) -> None:
    """Exact canonical DocumentVersionResponse shape; no SQLAlchemy
    internals leak through."""

    logical_document_id = uuid4()
    failed_version = make_document(
        logical_document_id=logical_document_id,
        status=DocumentStatus.FAILED,
    )
    ready_version = make_document(
        document_id=failed_version.id,
        logical_document_id=logical_document_id,
        status=DocumentStatus.READY,
        is_current=True,
    )

    ingestion = FakeIngestionService()
    ingestion.retry_result = ready_version

    override(get_document_repository, FakeDocumentRepository([failed_version]))
    override(get_ingestion_service, ingestion)

    response = client.post(f"/documents/{logical_document_id}/versions/{failed_version.id}/retry")

    assert response.status_code == 200

    body = response.json()
    assert set(body.keys()) == EXPECTED_VERSION_FIELDS
    assert FORBIDDEN_FIELDS.isdisjoint(body.keys())


def test_retry_requires_no_request_body(client: TestClient) -> None:
    """No upload body is required -- a bare POST with no data succeeds."""

    logical_document_id = uuid4()
    failed_version = make_document(
        logical_document_id=logical_document_id,
        status=DocumentStatus.FAILED,
    )
    ready_version = make_document(
        document_id=failed_version.id,
        logical_document_id=logical_document_id,
        status=DocumentStatus.READY,
        is_current=True,
    )

    ingestion = FakeIngestionService()
    ingestion.retry_result = ready_version

    override(get_document_repository, FakeDocumentRepository([failed_version]))
    override(get_ingestion_service, ingestion)

    # Deliberately no `data=`/`json=`/`files=` argument.
    response = client.post(f"/documents/{logical_document_id}/versions/{failed_version.id}/retry")

    assert response.status_code == 200


# ---------------------------------------------------------------------------
# 11. Deleting one version
# ---------------------------------------------------------------------------


def test_delete_document_version_succeeds(client: TestClient) -> None:
    logical_document_id = uuid4()
    version = make_document(logical_document_id=logical_document_id, is_current=False)

    deletion_service = FakeDocumentDeletionService()

    override(get_document_repository, FakeDocumentRepository([version]))
    override(get_document_deletion_service, deletion_service)

    response = client.delete(f"/documents/{logical_document_id}/versions/{version.id}")

    assert response.status_code == 204
    assert deletion_service.deleted_versions == [version.id]


def test_delete_document_version_response_has_no_body(client: TestClient) -> None:
    logical_document_id = uuid4()
    version = make_document(logical_document_id=logical_document_id, is_current=False)

    override(get_document_repository, FakeDocumentRepository([version]))
    override(get_document_deletion_service, FakeDocumentDeletionService())

    response = client.delete(f"/documents/{logical_document_id}/versions/{version.id}")

    assert response.status_code == 204
    assert response.content == b""


def test_delete_failed_version_succeeds(client: TestClient) -> None:
    """READY, FAILED, and PROCESSING versions may all be deleted."""

    logical_document_id = uuid4()
    failed_version = make_document(
        logical_document_id=logical_document_id,
        status=DocumentStatus.FAILED,
        is_current=False,
    )

    deletion_service = FakeDocumentDeletionService()

    override(get_document_repository, FakeDocumentRepository([failed_version]))
    override(get_document_deletion_service, deletion_service)

    response = client.delete(f"/documents/{logical_document_id}/versions/{failed_version.id}")

    assert response.status_code == 204
    assert deletion_service.deleted_versions == [failed_version.id]


def test_delete_processing_version_succeeds(client: TestClient) -> None:
    logical_document_id = uuid4()
    processing_version = make_document(
        logical_document_id=logical_document_id,
        status=DocumentStatus.PROCESSING,
        is_current=False,
    )

    deletion_service = FakeDocumentDeletionService()

    override(get_document_repository, FakeDocumentRepository([processing_version]))
    override(get_document_deletion_service, deletion_service)

    response = client.delete(f"/documents/{logical_document_id}/versions/{processing_version.id}")

    assert response.status_code == 204
    assert deletion_service.deleted_versions == [processing_version.id]


def test_delete_version_mismatched_with_logical_document_returns_404(
    client: TestClient,
) -> None:
    """version_id exists, but not under the logical_document_id in the URL --
    this must not reveal that the version exists elsewhere."""

    version = make_document(logical_document_id=uuid4(), is_current=False)

    deletion_service = FakeDocumentDeletionService()

    override(get_document_repository, FakeDocumentRepository([version]))
    override(get_document_deletion_service, deletion_service)

    other_logical_document_id = uuid4()
    response = client.delete(f"/documents/{other_logical_document_id}/versions/{version.id}")

    assert response.status_code == 404
    assert deletion_service.deleted_versions == []


def test_delete_version_both_ids_unknown_returns_404(client: TestClient) -> None:
    deletion_service = FakeDocumentDeletionService()

    override(get_document_repository, FakeDocumentRepository([]))
    override(get_document_deletion_service, deletion_service)

    response = client.delete(f"/documents/{uuid4()}/versions/{uuid4()}")

    assert response.status_code == 404
    assert deletion_service.deleted_versions == []


def test_delete_version_physical_file_failure_propagates(client: TestClient) -> None:
    """Physical-file deletion failure must propagate rather than pretending
    the deletion fully succeeded -- it must not be reported as 204."""

    logical_document_id = uuid4()
    version = make_document(logical_document_id=logical_document_id, is_current=False)

    deletion_service = FakeDocumentDeletionService()
    deletion_service.raise_on_delete_version = RuntimeError("Could not delete physical file.")

    override(get_document_repository, FakeDocumentRepository([version]))
    override(get_document_deletion_service, deletion_service)

    with pytest.raises(RuntimeError, match="Could not delete physical file."):
        client.delete(f"/documents/{logical_document_id}/versions/{version.id}")


# ---------------------------------------------------------------------------
# 12. Deleting an entire logical document
# ---------------------------------------------------------------------------


def test_delete_logical_document_succeeds(client: TestClient) -> None:
    """A logical document with a single version can be deleted."""

    logical_document_id = uuid4()
    version = make_document(logical_document_id=logical_document_id)

    deletion_service = FakeDocumentDeletionService()

    override(get_document_repository, FakeDocumentRepository([version]))
    override(get_document_deletion_service, deletion_service)

    response = client.delete(f"/documents/{logical_document_id}")

    assert response.status_code == 204
    assert deletion_service.deleted_logical_documents == [logical_document_id]


def test_delete_logical_document_with_multiple_versions_in_different_statuses_succeeds(
    client: TestClient,
) -> None:
    """The route does not care how many versions exist or what status each
    is in -- existence of the logical document is enough; the real removal
    of every version regardless of status is covered by
    DocumentDeletionService's own tests."""

    logical_document_id = uuid4()
    ready_version = make_document(
        logical_document_id=logical_document_id,
        version_number=1,
        status=DocumentStatus.READY,
        is_current=True,
    )
    failed_version = make_document(
        logical_document_id=logical_document_id,
        version_number=2,
        status=DocumentStatus.FAILED,
        is_current=False,
    )
    processing_version = make_document(
        logical_document_id=logical_document_id,
        version_number=3,
        status=DocumentStatus.PROCESSING,
        is_current=False,
    )

    deletion_service = FakeDocumentDeletionService()

    override(
        get_document_repository,
        FakeDocumentRepository([ready_version, failed_version, processing_version]),
    )
    override(get_document_deletion_service, deletion_service)

    response = client.delete(f"/documents/{logical_document_id}")

    assert response.status_code == 204
    assert deletion_service.deleted_logical_documents == [logical_document_id]


def test_delete_logical_document_response_has_no_body(client: TestClient) -> None:
    logical_document_id = uuid4()
    version = make_document(logical_document_id=logical_document_id)

    override(get_document_repository, FakeDocumentRepository([version]))
    override(get_document_deletion_service, FakeDocumentDeletionService())

    response = client.delete(f"/documents/{logical_document_id}")

    assert response.status_code == 204
    assert response.content == b""


def test_delete_logical_document_physical_file_failure_propagates(client: TestClient) -> None:
    """Physical-file deletion failure must propagate rather than pretending
    the deletion fully succeeded -- it must not be reported as 204."""

    logical_document_id = uuid4()
    version = make_document(logical_document_id=logical_document_id)

    deletion_service = FakeDocumentDeletionService()
    deletion_service.raise_on_delete_logical = RuntimeError("Could not delete physical file.")

    override(get_document_repository, FakeDocumentRepository([version]))
    override(get_document_deletion_service, deletion_service)

    with pytest.raises(RuntimeError, match="Could not delete physical file."):
        client.delete(f"/documents/{logical_document_id}")


# ---------------------------------------------------------------------------
# 13. Unknown document/version returns appropriate 404
# ---------------------------------------------------------------------------


def test_get_unknown_document_returns_404(client: TestClient) -> None:
    override(get_document_repository, FakeDocumentRepository([]))

    response = client.get(f"/documents/{uuid4()}")

    assert response.status_code == 404


def test_get_versions_of_unknown_document_returns_404(client: TestClient) -> None:
    override(get_document_repository, FakeDocumentRepository([]))

    response = client.get(f"/documents/{uuid4()}/versions")

    assert response.status_code == 404


def test_retry_unknown_version_returns_404(client: TestClient) -> None:
    ingestion = FakeIngestionService()

    override(get_document_repository, FakeDocumentRepository([]))
    override(get_ingestion_service, ingestion)

    response = client.post(f"/documents/{uuid4()}/versions/{uuid4()}/retry")

    assert response.status_code == 404
    assert ingestion.retry_calls == []


def test_retry_version_mismatched_with_logical_document_returns_404(
    client: TestClient,
) -> None:
    """version_id exists, but not under the logical_document_id in the URL."""

    version = make_document(logical_document_id=uuid4(), status=DocumentStatus.FAILED)
    ingestion = FakeIngestionService()

    override(get_document_repository, FakeDocumentRepository([version]))
    override(get_ingestion_service, ingestion)

    other_logical_document_id = uuid4()
    response = client.post(f"/documents/{other_logical_document_id}/versions/{version.id}/retry")

    assert response.status_code == 404
    assert ingestion.retry_calls == []


def test_delete_unknown_version_returns_404(client: TestClient) -> None:
    logical_document_id = uuid4()
    version = make_document(logical_document_id=logical_document_id)

    deletion_service = FakeDocumentDeletionService()

    override(get_document_repository, FakeDocumentRepository([version]))
    override(get_document_deletion_service, deletion_service)

    # version_id does not belong to this logical document at all.
    response = client.delete(f"/documents/{logical_document_id}/versions/{uuid4()}")

    assert response.status_code == 404
    assert deletion_service.deleted_versions == []


def test_delete_unknown_logical_document_returns_404(client: TestClient) -> None:
    deletion_service = FakeDocumentDeletionService()

    override(get_document_repository, FakeDocumentRepository([]))
    override(get_document_deletion_service, deletion_service)

    response = client.delete(f"/documents/{uuid4()}")

    assert response.status_code == 404
    assert deletion_service.deleted_logical_documents == []


# ---------------------------------------------------------------------------
# 14. API response shape does not expose SQLAlchemy internals
# ---------------------------------------------------------------------------


def test_response_shape_exposes_only_the_explicit_contract(client: TestClient) -> None:
    ingestion = FakeIngestionService()
    document = make_document()
    ingestion.next_document = document

    override(get_ingestion_service, ingestion)

    response = upload_file(client)

    assert response.status_code == 201

    body = response.json()

    # Exact key set: no SQLAlchemy relationship/internal attributes
    # (pages, chunks, sections, _sa_instance_state, raw metadata, ...)
    # leak through, and nothing from the explicit contract is missing.
    assert set(body.keys()) == EXPECTED_VERSION_FIELDS
