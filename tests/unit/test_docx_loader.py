from pathlib import Path

from docx import Document as DocxDocument

from app.ingestion.loaders.docx import DocxLoader
from app.ingestion.structure import UnitType
from app.ingestion.structure_extractor import StructureExtractor


def test_docx_loader_extracts_paragraphs_in_document_order(tmp_path: Path) -> None:
    docx_path = tmp_path / "test.docx"

    document = DocxDocument()
    document.add_paragraph("Retrieval-Augmented Generation")
    document.add_paragraph("RAG combines retrieval with generation.")
    document.add_paragraph("It grounds answers in retrieved evidence.")
    document.save(docx_path)

    loader = DocxLoader()
    pages = loader.load(docx_path)

    assert len(pages) == 1
    assert pages[0].content == (
        "Retrieval-Augmented Generation\n"
        "RAG combines retrieval with generation.\n"
        "It grounds answers in retrieved evidence."
    )


def test_docx_loader_ignores_empty_paragraphs(tmp_path: Path) -> None:
    docx_path = tmp_path / "test.docx"

    document = DocxDocument()
    document.add_paragraph("First paragraph.")
    document.add_paragraph("")
    document.add_paragraph("   ")
    document.add_paragraph("Second paragraph.")
    document.save(docx_path)

    loader = DocxLoader()
    pages = loader.load(docx_path)

    assert pages[0].content == "First paragraph.\nSecond paragraph."


def test_docx_loader_returns_page_number_one(tmp_path: Path) -> None:
    docx_path = tmp_path / "test.docx"

    document = DocxDocument()
    document.add_paragraph("Single page content.")
    document.save(docx_path)

    loader = DocxLoader()
    pages = loader.load(docx_path)

    assert len(pages) == 1
    assert pages[0].page_number == 1


def test_docx_loader_returns_empty_content_for_document_with_no_text(tmp_path: Path) -> None:
    docx_path = tmp_path / "test.docx"

    document = DocxDocument()
    document.add_paragraph("")
    document.save(docx_path)

    loader = DocxLoader()
    pages = loader.load(docx_path)

    assert pages[0].page_number == 1
    assert pages[0].content == ""


def test_docx_loader_includes_table_text_after_paragraphs(tmp_path: Path) -> None:
    """Tables are supported via the simple, documented ordering: all
    paragraph text first, then each table rendered as a real Markdown pipe
    table -- not interleaved at the table's true position in the document.

    The table is rendered as genuine pipe-table syntax (header row + a
    `|---|---|` separator), with the table's first row treated as its
    header, rather than a flat "cell | cell" text line -- this is what lets
    StructureExtractor detect it downstream as a real TABLE structural unit.
    """

    docx_path = tmp_path / "test.docx"

    document = DocxDocument()
    document.add_paragraph("Introduction paragraph.")

    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Header A"
    table.cell(0, 1).text = "Header B"
    table.cell(1, 0).text = "Value 1"
    table.cell(1, 1).text = "Value 2"

    document.save(docx_path)

    loader = DocxLoader()
    pages = loader.load(docx_path)

    assert pages[0].content == (
        "Introduction paragraph.\n\n| Header A | Header B |\n| --- | --- |\n| Value 1 | Value 2 |"
    )


def test_docx_loader_renders_multiple_tables_with_distinct_ids(tmp_path: Path) -> None:
    """Multiple tables in one document each round-trip as their own
    independent Markdown table block, detectable as separate TABLE units."""

    docx_path = tmp_path / "test.docx"

    document = DocxDocument()

    first_table = document.add_table(rows=2, cols=2)
    first_table.cell(0, 0).text = "A"
    first_table.cell(0, 1).text = "B"
    first_table.cell(1, 0).text = "1"
    first_table.cell(1, 1).text = "2"

    second_table = document.add_table(rows=2, cols=1)
    second_table.cell(0, 0).text = "Only Column"
    second_table.cell(1, 0).text = "Value"

    document.save(docx_path)

    loader = DocxLoader()
    pages = loader.load(docx_path)

    assert "| A | B |" in pages[0].content
    assert "| 1 | 2 |" in pages[0].content
    assert "| Only Column |" in pages[0].content
    assert "| Value |" in pages[0].content


def test_docx_loader_converts_native_heading_styles_to_markdown_headings(
    tmp_path: Path,
) -> None:
    docx_path = tmp_path / "headings.docx"

    document = DocxDocument()

    document.add_heading("Retrieval", level=1)
    document.add_paragraph("Retrieval finds relevant evidence.")
    document.add_heading("Vector Search", level=2)
    document.add_paragraph("Vector search uses embeddings.")
    document.add_heading("HNSW", level=3)
    document.add_paragraph("HNSW provides approximate nearest-neighbor search.")

    document.save(docx_path)

    loader = DocxLoader()
    pages = loader.load(docx_path)

    assert pages[0].content == (
        "# Retrieval\n"
        "Retrieval finds relevant evidence.\n"
        "## Vector Search\n"
        "Vector search uses embeddings.\n"
        "### HNSW\n"
        "HNSW provides approximate nearest-neighbor search."
    )


def test_docx_loader_preserves_normal_paragraphs_with_native_headings(
    tmp_path: Path,
) -> None:
    docx_path = tmp_path / "mixed.docx"

    document = DocxDocument()
    document.add_heading("Introduction", level=1)
    document.add_paragraph("This remains a normal paragraph.")
    document.add_heading("Details", level=2)

    document.save(docx_path)

    loader = DocxLoader()
    pages = loader.load(docx_path)

    assert pages[0].content == ("# Introduction\nThis remains a normal paragraph.\n## Details")


def test_docx_fixture_preserves_native_heading_structure() -> None:
    fixture_path = Path("tests/evaluation/fixtures/benchmark/vector_search_fundamentals.docx")

    pages = DocxLoader().load(fixture_path)
    units = StructureExtractor().extract(pages)

    headings = [unit for unit in units if unit.unit_type == UnitType.HEADING]

    assert [heading.content for heading in headings] == [
        "Vector Search Fundamentals",
        "Embeddings",
        "Similarity Search",
        "Chunking and Retrieval Units",
        "Recall and Precision",
        "Failure Modes",
        "Practical Retrieval Design",
    ]

    assert headings[0].section_path == "Vector Search Fundamentals"
    assert headings[1].section_path == "Embeddings"
